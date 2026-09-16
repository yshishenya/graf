import json
from copy import deepcopy
from dataclasses import replace
from io import BytesIO

import pytest
from openpyxl import load_workbook

from tests.unit.test_meeting_protocol import protocol_fixture, validate
from tests.unit.test_transcript_exports import _snapshot
from twobrain_rec_server.cabinet.access import narrow_summary_projection
from twobrain_rec_server.cabinet.exports import render_content_export
from twobrain_rec_server.cabinet.meeting_protocol import (
    TASK_HEADING,
    protocol_lines,
    without_protocol_evidence,
)
from twobrain_rec_server.cabinet.rendering import _render_full_protocol


@pytest.mark.parametrize("version,label", [
    ("meeting-protocol-v1-about-v1", "О чём встреча"),
    ("meeting-protocol-v1", "Тип встречи"),
    ("unknown", "Тип встречи"),
    (None, "Тип встречи"),
])
def test_detail_and_shared_identity_use_current_metadata(version, label):
    from datetime import UTC, datetime

    from tests.unit.test_cabinet_web_shell import _review
    from twobrain_rec_server.api.schemas import OutcomeProvenanceView
    from twobrain_rec_server.cabinet.rendering import (
        render_meeting_detail_page,
        render_shared_meeting_summary_page,
    )

    review = _review()
    document = validate(protocol_fixture())["protocol"]
    document.update(title="Устаревшее название протокола", date_and_time="Придуманное время",
                    meeting_type="Техническое проектирование уведомлений", participants=["<Участник>"])
    review.meeting.title = "Актуальное название встречи"
    review.meeting.started_at = datetime(2026, 9, 8, 22, 5, 39, tzinfo=UTC)
    review.notes_action_truth.protocol = document
    if version is not None:
        review.notes_action_truth.provenance = OutcomeProvenanceView(
            generator_kind="litellm", generator_version=version,
        )
    detail = render_meeting_detail_page(review)
    shared = render_shared_meeting_summary_page(
        meeting_title=review.meeting.title, occurred_at=review.meeting.started_at,
        duration_seconds=120, summary_sections=[], protocol=document, generator_version=version,
    )
    for html in (detail, shared):
        assert html.split("<template")[0].count("<h1") == 1
        assert label in html
        assert "Техническое проектирование уведомлений" in html
        assert "&lt;Участник&gt;" in html
        assert "Устаревшее название протокола" not in html
        assert "Придуманное время" not in html
        assert "2026-09-08T22:05:39" in html
    assert "data-source-segment" not in shared


def test_interactive_order_disclosures_and_complete_export_are_independent():
    document = validate(protocol_fixture())["protocol"]
    document["notes"] = [{"text": "Синтетическое примечание", "source_refs": []}]
    document["topics"][0]["outcome"] = [{"text": "Выбор отложен", "source_refs": []}]
    document["topics"][0]["context"] = [{"text": "Нужен доступ к макету", "source_refs": []}]
    document["topics"][0]["proposals"] = []
    html = _render_full_protocol(document, source_destination_available=False)
    titles = ["Главное", "Принятые решения", "Задачи", "Открытые вопросы и следующие шаги",
              "Ключевые обсуждения", "Примечания"]
    offsets = [html.index(f">{title}<") for title in titles]
    assert offsets == sorted(offsets)
    assert '<details class="notes-topic">' in html
    assert '<details class="notes-section notes-protocol-notes">' in html
    assert ' open' not in html
    assert html.index(">Итог<") < html.index(">Контекст<") < html.index(">Обсуждение<")
    assert ">Предложения<" not in html
    assert "Цели встречи" not in html and "Тип входных данных" not in html
    assert "Выбор отложен" in html and "Синтетическое примечание" in html
    exported = "\n".join(protocol_lines(document, markdown=False, include_evidence=False, meeting_url=""))
    assert "Цели встречи" in exported and "Тип входных данных" in exported
    assert exported.index("Контекст") < exported.index("Итог\n")
    document["action_items"] = []
    document["notes"] = []
    html = _render_full_protocol(document, source_destination_available=False)
    assert "Задачи не зафиксированы" in html and "<table" not in html
    assert ">Примечания<" not in html


@pytest.mark.parametrize("count", [0, 1, 4])
def test_one_visible_source_and_no_empty_overflow(count):
    from twobrain_rec_server.api.schemas import OutcomeItemView, OutcomeSourceReferenceView
    from twobrain_rec_server.cabinet.rendering import _render_outcome_item

    item = OutcomeItemView(category="summary", sequence=0, text="Проверяемый тезис", truth_label="supported",
        owner_text="Анна",
        source_refs=[OutcomeSourceReferenceView(
            transcript_segment_id="00000000-0000-0000-0000-000000000002",
            start_seconds=seconds, evidence_kind="segment", seekable=True,
        ) for seconds in range(count)])
    html = _render_outcome_item(item, source_destination_available=True)
    assert html.split("<details")[0].count("data-seek-seconds") == min(count, 1)
    assert html.count("data-seek-seconds") == count
    assert ("notes-source-more" in html) == (count > 1)
    if count:
        assert html.index('notes-item-sources') < html.index('notes-item-meta-row')
    if count > 1:
        assert f"Ещё {count - 1}" in html
        assert '<div class="notes-source-list">' in html


def test_full_protocol_rendering_has_sections_task_table_and_canonical_links():
    document = validate(protocol_fixture())["protocol"]
    text = "\n".join(protocol_lines(document, markdown=True, include_evidence=True,
                                    meeting_url="https://graf.example/cabinet/meetings/synthetic"))
    assert "## Ключевые обсуждения" in text
    assert "### Макет" in text
    assert "| Задача | Ответственный | Срок |" in text
    assert "наверное, завтра" in text
    assert "[00:12](https://graf.example/cabinet/meetings/synthetic?source_result_id=00000000-0000-0000-0000-000000000001#graf-source=" in text
    assert "макет\\," not in text
    html = _render_full_protocol(document, source_destination_available=True,
        source_targets=frozenset({("00000000-0000-0000-0000-000000000001", "00000000-0000-0000-0000-000000000002")}))
    assert '<table' in html and '<th role="columnheader" scope="col">Ответственный</th>' in html
    assert 'data-seek-seconds="12.25"' in html
    assert 'data-source-segment="00000000-0000-0000-0000-000000000002"' in html
    assert "Риски не найдены" not in html


def test_protocol_escapes_unsafe_output_and_does_not_mutate_evidence():
    document = validate(protocol_fixture())["protocol"]
    document["title"] = '<script>alert(1)</script> ![image](https://evil.example)'
    original = deepcopy(document)
    html = _render_full_protocol(document, source_destination_available=False)
    assert "<script>" not in html and "<img" not in html
    assert "data-seek-seconds" not in html
    text = "\n".join(protocol_lines(document, markdown=True, include_evidence=False, meeting_url=""))
    assert "![image]" not in text and "#graf-source" not in text
    stripped = without_protocol_evidence(document)
    assert "source_refs" not in str(stripped) and "quote" not in str(stripped)
    assert document == original


def test_empty_states_and_unknown_task_fields():
    document = protocol_fixture()
    document["action_items"] = []
    text = "\n".join(protocol_lines(validate(document)["protocol"], markdown=True,
                                    include_evidence=False, meeting_url=""))
    assert "Принятые решения в расшифровке не зафиксированы." in text
    assert "| Задачи не зафиксированы | Не назначен | Не указан |" in text
    assert "Открытые вопросы не зафиксированы." in text


@pytest.mark.parametrize("scope", ["summary", "combined"])
def test_saved_protocol_exports_preserve_words_and_hide_evidence(export_fixture, scope):
    document = validate(protocol_fixture())["protocol"]
    document["title"] = "=SUM(1,2)"
    for format in ("md", "txt", "json", "xlsx"):
        snapshot = _snapshot(export_fixture, scope=scope, format=format)
        snapshot = replace(snapshot, summary=replace(snapshot.summary, protocol=document),
                           selection=replace(snapshot.selection, include_evidence=False))
        exported = render_content_export(snapshot).body
        if format == "json":
            parsed = json.loads(exported)
            assert "source_refs" not in json.dumps(parsed["summary"]["protocol"])
            assert parsed["summary"]["protocol"]["title"] == document["title"]
        elif format == "xlsx":
            workbook = load_workbook(BytesIO(exported))
            assert workbook.sheetnames == ["Transcript", "Summary", "Action Items", "Metadata", "Протокол"]
            assert dict(workbook["Metadata"].values)["schema_version"] == "graf.transcript-export.v3"
            sheet = workbook["Протокол"]
            assert all(cell.data_type != "f" for row in sheet for cell in row)
            values = list(sheet.values)
            assert any(row[1] == "Принятые решения в расшифровке не зафиксированы." for row in values)
            assert any(row[1] == "Сделать макет" and row[2] == "Участник 1" for row in values)
            assert all(not row[4] for row in values[1:])
            assert all(cell.hyperlink is None for row in sheet for cell in row)
        else:
            text = exported.decode()
            assert "Сделать макет" in text and "наверное, завтра" in text
            assert "Источник сохраненной ревизии" not in text and "graf-source" not in text
    assert document["action_items"][0]["source_refs"]
    snapshot = _snapshot(export_fixture, scope="summary", format="xlsx")
    snapshot = replace(snapshot, summary=replace(snapshot.summary, protocol=document))
    sheet = load_workbook(BytesIO(render_content_export(snapshot).body))["Протокол"]
    links = [cell.hyperlink for row in sheet for cell in row if cell.hyperlink]
    assert links and all("#graf-source=00000000-0000-0000-0000-000000000002" in link.target for link in links)


def test_summary_only_share_has_complete_protocol_without_transcript_refs():
    from datetime import UTC, datetime
    document = validate(protocol_fixture())["protocol"]
    shared = narrow_summary_projection(meeting_label="Встреча", occurred_at=datetime.now(UTC),
                                       duration_seconds=30, summary_sections=[], protocol=document)
    assert shared["protocol"]["topics"][0]["title"] == "Макет"
    assert "source_refs" not in json.dumps(shared["protocol"]) and "quote" not in json.dumps(shared["protocol"])
    html = _render_full_protocol(shared["protocol"], source_destination_available=False)
    assert "Сделать макет" in html and "data-source-segment" not in html


@pytest.mark.parametrize("field,column", [("task", 1), ("owner_text", 2), ("due_date_text", 3)])
def test_xlsx_preserves_long_protocol_fields_in_continuation_rows(export_fixture, field, column):
    document = validate(protocol_fixture())["protocol"]
    # The second chunk starts with a formula marker; it must remain literal.
    value = "я" * 32767 + "=SUM(1,2)\n" + "🙂" * 33000 + " конец"
    document["action_items"][0][field] = value
    snapshot = _snapshot(export_fixture, scope="summary", format="xlsx")
    snapshot = replace(snapshot, summary=replace(snapshot.summary, protocol=document),
                       selection=replace(snapshot.selection, include_evidence=False))
    sheet = load_workbook(BytesIO(render_content_export(snapshot).body))["Протокол"]
    rows = list(sheet.values)
    start = next(index for index, row in enumerate(rows) if row[0] == TASK_HEADING) + 1
    parts = []
    for row in rows[start:]:
        if row[0]:
            break
        if row[column] is not None:
            parts.append(row[column])
    assert "".join(parts) == value
    assert len(parts) >= 3 and all(len(part.encode("utf-16-le")) <= 32767 * 2 for part in parts)
    assert all(cell.data_type != "f" for row in sheet for cell in row)
    assert document["action_items"][0][field] == value


def test_protocol_source_links_pin_revision_in_all_hyperlink_exports(export_fixture):
    from twobrain_rec_server.cabinet.meeting_protocol import protocol_source_url

    document = validate(protocol_fixture())["protocol"]
    ref = document["action_items"][0]["source_refs"][0]
    url = protocol_source_url("https://graf.example/meetings/id?summary_format=auto#old", ref)
    assert "summary_format=auto&source_result_id=" + ref["processing_result_id"] in url
    assert url.endswith("#graf-source=" + ref["transcript_segment_id"])
    for format in ("md", "txt", "xlsx"):
        snapshot = _snapshot(export_fixture, scope="summary", format=format)
        snapshot = replace(snapshot, summary=replace(snapshot.summary, protocol=document),
                           selection=replace(snapshot.selection, include_evidence=True))
        body = render_content_export(snapshot).body
        if format == "xlsx":
            sheet = load_workbook(BytesIO(body))["Протокол"]
            links = [cell.hyperlink.target for row in sheet for cell in row if cell.hyperlink]
            assert links and all("source_result_id=" in link for link in links)
        else:
            assert "?source_result_id=" in body.decode()


@pytest.mark.parametrize("result_id,segment_id", [
    ("other-revision", "00000000-0000-0000-0000-000000000002"),
    ("00000000-0000-0000-0000-000000000001", "missing-segment"),
])
def test_protocol_never_seeks_into_a_different_source(result_id, segment_id):
    document = validate(protocol_fixture())["protocol"]
    html = _render_full_protocol(document, source_destination_available=True,
                                source_targets=frozenset({(result_id, segment_id)}))
    assert "Сделать макет" in html
    assert "data-seek-seconds" not in html


def test_meeting_notes_bind_inline_sources_to_the_displayed_transcript():
    from types import SimpleNamespace

    from twobrain_rec_server.cabinet.rendering import _render_notes_outcomes
    document = validate(protocol_fixture())["protocol"]
    ref = document["action_items"][0]["source_refs"][0]
    segment = SimpleNamespace(processing_result_id=ref["processing_result_id"], segment_id=ref["transcript_segment_id"])
    review = SimpleNamespace(notes_action_truth=SimpleNamespace(protocol=document),
        transcript=SimpleNamespace(available=True, segments=[segment], speaker_turns=[]),
        playback=SimpleNamespace(can_play=True))
    assert 'data-seek-seconds="12.25"' in _render_notes_outcomes(review)
    segment.processing_result_id = "newer-result"
    assert "data-seek-seconds" not in _render_notes_outcomes(review)
